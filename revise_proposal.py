from docx import Document
from docx.oxml import OxmlElement

PATH = 'deliverables/Smart_Asset_Management_Revised.docx'

replacements = {
    'Machine Learning Enabled Smart Asset Management System for Predictive Maintenance, Asset Lifecycle Management and Google-Based Device Location Tracking':
        'Machine Learning Enabled Smart Asset Management System for Predictive Maintenance and Google-Based Device Location Tracking',
    'Machine learning features will support predictive maintenance, smart asset allocation, asset lifecycle prediction, anomaly detection and automated alerts. By analysing device age, repair history, warranty status, usage patterns, assignment records and permitted location events, the system will help administrators identify high-risk assets and make better decisions before problems become serious.':
        'The machine-learning component will focus on predictive maintenance: estimating whether an ICT asset is likely to require maintenance within the next 30 days. Rule-based security and location alerts will complement this prediction. By analysing device age, repair history, warranty status and usage patterns, the system will help administrators identify high-risk assets before problems become serious.',
    'This project, therefore, seeks to develop a Smart Asset Management System that automates asset tracking and management while integrating consent-based device-location tracking through Google Maps Platform, machine learning for predictive maintenance, smart asset allocation, anomaly detection, lifecycle prediction, and automated alerts.':
        'This project, therefore, seeks to develop a Smart Asset Management System that automates asset tracking and management while integrating consent-based device-location check-ins through Google Maps Platform, machine learning for predictive maintenance, and automated rule-based alerts.',
    'To design and develop a secure web-based Smart Asset Management System that enables organisations to register, assign, track, monitor, and manage ICT assets while applying machine learning techniques to support predictive maintenance, smart asset allocation, anomaly detection, and asset lifecycle management.':
        'To design and develop a secure web-based Smart Asset Management System that enables organisations to register, assign, track, monitor and manage ICT assets while applying machine learning to support predictive maintenance.',
    'To apply machine learning techniques for predictive maintenance, asset lifecycle prediction, smart allocation, and anomaly detection':
        'To apply a machine-learning model to predict whether an ICT asset is likely to require maintenance within the next 30 days',
    'How can machine learning be used to support predictive maintenance and asset lifecycle management?':
        'How can machine learning be used to identify ICT assets likely to require maintenance within the next 30 days?',
    'It will incorporate predictive maintenance, asset lifecycle prediction, anomaly detection, smart asset allocation and a Google Maps-based device-location module.':
        'It will incorporate predictive maintenance and a Google Maps-based device-location module. Rule-based alerts will flag selected security and location events.',
    'The machine learning component will focus on selected prediction tasks – Only predictive maintenance, lifecycle prediction, anomaly detection, and smart allocation will be implemented.':
        'The machine-learning component will focus on one prediction task – predictive maintenance. Lifecycle forecasting and smart allocation are future enhancements; anomaly alerts are implemented as transparent rules rather than a separate machine-learning model.',
    'Historical and operational asset data will be analysed to support prediction, classification, and anomaly detection tasks.':
        'Historical and operational asset data will be analysed to support one binary classification task: predictive maintenance.',
    'Machine learning models will be trained to support predictive maintenance, lifecycle prediction, and anomaly detection. The models may use features such as asset age, repair frequency, usage period, status history, warranty status, and login behaviour. The most suitable model will be selected based on the type and quality of the available data.':
        'The predictive-maintenance model will be trained to classify whether an asset is likely to require maintenance within the next 30 days. It will use asset age, asset type, warranty remaining, repair count in the previous 12 months, days since last repair, unresolved-fault indicator, average daily usage hours, monthly crash count and battery-health percentage. Model selection will be based on validation performance.',
    'The project uses two machine learning algorithms to enhance ICT asset management. Random Forest, a supervised learning algorithm, is applied to predictive maintenance and asset lifecycle forecasting, selected for its strong prediction accuracy, ability to handle multiple asset-related variables and flexibility across both classification and regression tasks. Isolation Forest complements this by detecting anomalies such as unusual login activity, repeated failed access attempts, unexpected asset movements, irregular usage patterns and suspicious behaviour. These algorithms are implemented using Python libraries such as Scikit-learn, Pandas and NumPy, with Flask serving as the integration layer between the machine learning services and the PHP-based asset management system.':
        'The project uses Logistic Regression as the primary predictive-maintenance classifier because it produces an interpretable probability and shows how each feature is associated with maintenance risk. A Random Forest classifier will be trained only as a performance comparator for the same 30-day maintenance prediction; it is not a separate feature. The selected model will be the one that meets the stated validation criteria while remaining suitable for explanation to administrators. Isolation Forest, lifecycle prediction and smart allocation are outside this project’s implemented ML scope. Transparent rule-based alerts will instead flag defined security and location events. Python libraries such as scikit-learn, Pandas and NumPy will be used, with Flask serving as the integration layer between the model and the PHP-based system.'
}

doc = Document(PATH)
for p in doc.paragraphs:
    text = p.text
    if text in replacements:
        p.text = replacements[text]

# A few proposal paragraphs contain minor wording variations; update them by
# their stable opening wording while retaining all unrelated paragraphs.
prefix_replacements = {
    'The project uses two machine learning algorithms to enhance ICT asset management.':
        'The project uses Logistic Regression as the primary predictive-maintenance classifier because it produces an interpretable probability and shows how each feature is associated with maintenance risk. A Random Forest classifier will be trained only as a performance comparator for the same 30-day maintenance prediction; it is not a separate feature. The selected model will be the one that meets the stated validation criteria while remaining suitable for explanation to administrators. Isolation Forest, lifecycle prediction and smart allocation are outside this project’s implemented ML scope. Transparent rule-based alerts will instead flag defined security and location events. Python libraries such as scikit-learn, Pandas and NumPy will be used, with Flask serving as the integration layer between the model and the PHP-based system.',
}
for p in doc.paragraphs:
    for prefix, replacement in prefix_replacements.items():
        if p.text.startswith(prefix):
            p.text = replacement
            break

anchor_text = 'The project uses Logistic Regression as the primary predictive-maintenance classifier'
anchor = next(p for p in doc.paragraphs if p.text.startswith(anchor_text))

sections = [
    ('Heading 2', 'Machine Learning Scope, Dataset and Evaluation Clarification'),
    ('Heading 3', 'Primary Machine-Learning Task'),
    ('Normal', 'The project implements one machine-learning function: predictive maintenance. For each ICT asset, the model predicts the binary target maintenance_required_within_30_days (1 = a maintenance record or verified fault is expected within the next 30 days; 0 = no such event). The output is a probability and a Low, Medium or High risk label to help an administrator prioritise inspection, repair or replacement. Asset lifecycle prediction and smart allocation will not be implemented as ML models in this project; they remain future enhancements. Security and location anomalies are handled through explicit rules, not a second ML task.'),
    ('Heading 3', 'Dataset and Target Variable'),
    ('Normal', 'The training dataset will be a reproducible synthetic ICT-asset maintenance dataset created for this academic prototype because sufficiently complete, privacy-cleared organisational maintenance histories are not available. Each row represents one asset at a defined monthly observation date. Its schema mirrors the system database and maintenance register: asset type, age in months, remaining warranty days, repair count in the previous 12 months, days since last repair, unresolved-fault indicator, average daily usage hours, crash count in the previous month and battery-health percentage where applicable.'),
    ('Normal', 'Synthetic records will be generated using documented, realistic ranges and relationships: older assets, repeated repairs, unresolved faults, high crash counts, poor battery health and expired warranties will have a higher probability of a maintenance event. The 30-day target will then be assigned by the same simulated future maintenance event, not copied from an input feature. Generation rules, distributions, random seed and assumptions will be documented so the dataset can be reproduced and audited. The synthetic data supports a proof of concept only; the proposal will not claim that results generalise to a real organisation until the model is retrained and externally validated on consented operational data.'),
    ('Heading 3', 'Models and Baseline'),
    ('Normal', 'Logistic Regression is the primary model because it is simple, explainable and appropriate for a binary maintenance-risk outcome. Random Forest will be used only as a comparator on exactly the same dataset and features, to determine whether nonlinear relationships yield a material improvement. A most-frequent-class DummyClassifier is the minimum baseline. The final deployed choice will be justified by validation performance and interpretability; no lifecycle, allocation or anomaly-detection model will be trained.'),
    ('Heading 3', 'Data Division and Model Evaluation'),
    ('Normal', 'After stratifying by the maintenance target, the dataset will be divided into 70% training, 15% validation and 15% test data. Pre-processing steps and model fitting will be performed using the training set only. The validation set will select model hyperparameters and the probability threshold; the test set will be used once for the final unbiased evaluation. Where the generated dataset is small, five-fold cross-validation on the training portion will supplement, not replace, the held-out validation and test sets.'),
    ('Normal', 'Because missing a genuinely at-risk asset is costly, recall for the maintenance-required class is the primary metric. Precision, F1-score, confusion matrix, PR-AUC and ROC-AUC will also be reported. A model will be considered acceptable for the prototype only if, on the held-out test data, it achieves recall of at least 0.70, precision of at least 0.60, F1-score of at least 0.65 and PR-AUC at least 0.10 higher than the DummyClassifier baseline. These are acceptance criteria, not claimed results. Logistic Regression and Random Forest will be compared using the same test split; if their performance is similar, Logistic Regression will be selected for clearer explanations.'),
    ('Heading 3', 'Rule-Based Anomaly Alerts'),
    ('Normal', 'In this system an anomaly is a defined event that departs from approved policy or an asset’s expected record, rather than an ML-generated outlier. Examples are five or more failed login attempts from one account within 15 minutes, a status or assignment change by an unauthorised role, an asset assignment with no recorded handover, or an authorised location check-in outside its approved area. Each alert records the rule triggered, time, related user or asset and acknowledgement status. The rules will be evaluated by injecting normal and abnormal test events and reporting detection rate, false-alert rate and alert-generation time.'),
    ('Heading 3', 'Asset Lifecycle and Smart Allocation Boundaries'),
    ('Normal', 'For this project, “asset lifecycle” remains a descriptive status field (for example: active, under maintenance, retired), not a predictive model. “Smart allocation” will not make automated recommendations; normal allocation remains a rule- and administrator-led process based on availability, compatibility, department and user need. Their prediction/recommendation accuracy is therefore not evaluated as ML results. This boundary keeps the implementation and evaluation aligned with the single approved ML task.'),
    ('Heading 3', 'Google-Based Device Location Tracking'),
    ('Normal', 'Google Maps does not track a device. It only displays coordinates supplied to the application. In the prototype, the device user initiates a browser location check-in over HTTPS and grants the browser Geolocation API permission; the application stores the resulting last-known latitude, longitude, accuracy and timestamp, then displays that record on Google Maps. The system cannot use Google Maps or Google Find My Device to remotely locate an arbitrary laptop or phone, and it does not perform covert or continuous tracking.'),
    ('Heading 3', 'System Effectiveness Evaluation'),
    ('Normal', 'System evaluation will combine functional tests with a small, consented usability and task-based evaluation involving intended users such as ICT administrators. The same representative asset-management tasks will be completed using the prototype and the current spreadsheet/manual process where available. Measures will include: task completion time for registering, locating and reporting an asset; task success rate; record completeness (percentage of mandatory fields completed); discrepancy rate between a sampled physical asset list and database records; time taken to identify overdue maintenance; and the number and resolution time of rule-based alerts. For the location module, the study will record successful consented check-ins, timestamp freshness, reported accuracy and approved-area alert correctness.'),
    ('Normal', 'Participants will complete a short System Usability Scale (SUS) questionnaire and provide feedback on the usefulness of maintenance-risk explanations. A SUS score of 68 or above, a lower median task-completion time and fewer record discrepancies than the comparison process will indicate improved usability and asset-management effectiveness. Predictive-maintenance effectiveness will be evaluated separately against the held-out labelled test set using the stated ML metrics; it will not be inferred from user opinion alone.')
]

new_paras = []
for style, text in sections:
    p = doc.add_paragraph(text, style=style)
    new_paras.append(p)

# Place the clarification directly after the ML tools subsection.
cursor = anchor._p
for p in new_paras:
    cursor.addnext(p._p)
    cursor = p._p

doc.save(PATH)
