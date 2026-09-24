from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).parent
OUT = ROOT / 'deliverables' / 'Smart_Asset_Management_Final_Documentation.docx'
FIG = ROOT / 'output' / 'final_documentation_figures'
FIG.mkdir(parents=True, exist_ok=True)

def f(size, bold=False):
    path = '/System/Library/Fonts/Supplemental/Arial Bold.ttf' if bold else '/System/Library/Fonts/Supplemental/Arial.ttf'
    return ImageFont.truetype(path, size) if Path(path).exists() else ImageFont.load_default()

def diagram(name, title, columns, lines):
    im = Image.new('RGB', (1600, 900), 'white'); d = ImageDraw.Draw(im)
    d.text((55, 35), title, fill='#173f5f', font=f(34, True))
    xs = [170 + i * (1260 // max(1, len(columns) - 1)) for i in range(len(columns))]
    for x, heading in zip(xs, columns):
        d.rounded_rectangle((x-135, 140, x+135, 230), 18, fill='#e8f1f8', outline='#173f5f', width=3)
        d.multiline_text((x-110, 165), heading, fill='#173f5f', font=f(20, True), spacing=4)
        d.line((x, 230, x, 760), fill='#9aa7b0', width=2)
    y = 285
    for source, target, label in lines:
        x1, x2 = xs[source], xs[target]
        d.line((x1, y, x2, y), fill='#1f77b4', width=4)
        direction = 16 if x2 >= x1 else -16
        d.polygon([(x2, y), (x2-direction, y-8), (x2-direction, y+8)], fill='#1f77b4')
        d.rounded_rectangle((min(x1,x2)+20, y-27, max(x1,x2)-20, y-5), 5, fill='white')
        d.text((min(x1,x2)+30, y-25), label, fill='#222222', font=f(16))
        y += 85
    im.save(FIG / name)
    return FIG / name

figures = [
    ('Figure 4.1: Use-case diagram', diagram('use_case.png', 'Use-case diagram', ['Administrator', 'System', 'Employee'], [(0,1,'register assets; review maintenance; assign laptops'), (2,1,'view assigned laptop and maintenance status'), (0,1,'manage users and reports')])),
    ('Figure 4.2: Entity-relationship diagram', diagram('erd.png', 'Entity-relationship diagram', ['Users', 'Laptops', 'Maintenance records', 'Lifecycle history'], [(0,1,'one user may receive laptops'), (1,2,'one laptop has maintenance records'), (1,3,'one laptop has lifecycle events')])),
    ('Figure 4.3: Class diagram', diagram('class.png', 'Class diagram', ['User', 'Laptop', 'MaintenanceRecord', 'Recommendation'], [(0,1,'assignedTo'), (1,2,'has records'), (3,1,'ranks available laptops')])),
    ('Figure 4.4: Activity diagram', diagram('activity.png', 'Needs-based assignment activity', ['Administrator', 'Recommendation model', 'Asset register'], [(0,1,'enter department, rank and work type'), (1,2,'rank available laptops by specifications'), (0,2,'confirm recommended assignment')])),
    ('Figure 4.5: Sequence diagram', diagram('sequence.png', 'Assignment sequence diagram', ['Administrator', 'PHP application', 'Database'], [(0,1,'submit employee needs'), (1,2,'read available laptops and specifications'), (1,0,'display ranked recommendations'), (0,1,'confirm assignment'), (1,2,'save assignment and lifecycle history')])),
    ('Figure 4.6: System architecture diagram', diagram('architecture.png', 'System architecture diagram', ['Web browser', 'PHP application', 'MySQL database', 'ML service'], [(0,1,'secure request'), (1,2,'store assets, history and maintenance'), (1,3,'maintenance-risk request'), (3,1,'risk score response')])),
]

doc = Document(); sec = doc.sections[0]; sec.top_margin = Inches(.8); sec.bottom_margin = Inches(.8)
styles = doc.styles; styles['Normal'].font.name='Arial'; styles['Normal'].font.size=Pt(10.5)
title=doc.add_paragraph(); title.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=title.add_run('SMART ASSET MANAGEMENT SYSTEM'); r.bold=True; r.font.size=Pt(22)
sub=doc.add_paragraph(); sub.alignment=WD_ALIGN_PARAGRAPH.CENTER; sub.add_run('Final Project Documentation\nPredictive Maintenance and Asset Lifecycle Management').font.size=Pt(14)
doc.add_paragraph('\nPrepared for the Bachelor of Science in Informatics and Computer Science Project').alignment=WD_ALIGN_PARAGRAPH.CENTER
doc.add_page_break()
for heading, body in [
    ('Chapter 1: Introduction', 'The system provides a central register for organisational laptops, their assignments, maintenance records and lifecycle status. Its intelligent feature is predictive maintenance: it estimates the likelihood that an asset will require maintenance within 30 days.'),
    ('Chapter 2: Requirements and Scope', 'Users register laptops, record specifications and lifecycle status, create maintenance records, view maintenance risk and assign laptops. Administrators use a needs-based recommendation workflow: department, organisational rank and primary work type determine the target specification, then available laptops are ranked. Lifecycle tracking records Available, Assigned, Maintenance, Retired and Disposed states.'),
    ('Chapter 3: Design and Implementation', 'The solution uses PHP, MySQL, HTML and CSS. MySQL stores users, laptops, maintenance records, risk predictions and laptop history. The assignment page uses transparent suitability scoring based on processor tier, RAM, storage and department match; the administrator remains responsible for the final assignment. Logistic Regression remains the primary predictive-maintenance model, while Random Forest may be compared during evaluation.'),
]:
    doc.add_heading(heading, level=1); doc.add_paragraph(body)
doc.add_heading('Chapter 4: System Design Diagrams', level=1)
doc.add_paragraph('The following diagrams document the implemented system and the needs-based assignment workflow.')
for caption, image in figures:
    doc.add_heading(caption, level=2); doc.add_picture(str(image), width=Inches(6.5)); p=doc.add_paragraph(caption); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
doc.add_heading('Chapter 5: Testing and Evaluation', level=1)
doc.add_paragraph('Functional testing verifies login, asset registration, maintenance recording, lifecycle-status changes, recommendation ranking and assignment. Evaluation measures include task completion time, record completeness, maintenance-risk recall, precision, F1-score and PR-AUC against a baseline classifier. Assignment recommendations are reviewed for suitability against the defined minimum specification; administrators confirm every assignment.')
doc.add_heading('Deployment Notes', level=1)
doc.add_paragraph('Run database.sql for a new installation. Existing installations must run migrations/20260902_needs_based_assignment.sql before using the needs-based assignment page. The project should be deployed with database credentials kept outside source control.')
doc.save(OUT)
print(OUT)
