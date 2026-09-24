from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from docx.enum.table import WD_ROW_HEIGHT_RULE
from pathlib import Path

OUT = Path('output/Smart_Asset_Management_Updated_Proposal.docx')
OUT.parent.mkdir(parents=True, exist_ok=True)

BLUE = '1F4D78'; MID = '2E74B5'; PALE = 'E8EEF5'; GREY = 'F2F4F7'; INK = '202020'

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr(); mar = tcPr.first_child_found_in('w:tcMar')
    if mar is None: mar = OxmlElement('w:tcMar'); tcPr.append(mar)
    for side, value in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node = mar.find(qn('w:'+side))
        if node is None: node=OxmlElement('w:'+side); mar.append(node)
        node.set(qn('w:w'), str(value)); node.set(qn('w:type'),'dxa')

def set_table_widths(table, widths):
    table.autofit = False
    tblPr = table._tbl.tblPr
    tblW = tblPr.first_child_found_in('w:tblW')
    if tblW is None: tblW=OxmlElement('w:tblW'); tblPr.append(tblW)
    tblW.set(qn('w:w'), '9360'); tblW.set(qn('w:type'),'dxa')
    ind=OxmlElement('w:tblInd'); ind.set(qn('w:w'),'120'); ind.set(qn('w:type'),'dxa'); tblPr.append(ind)
    grid = table._tbl.tblGrid
    for col, w in zip(grid.gridCol_lst, widths): col.set(qn('w:w'), str(w))
    for row in table.rows:
        for cell,w in zip(row.cells,widths):
            tcPr=cell._tc.get_or_add_tcPr(); tcW=tcPr.first_child_found_in('w:tcW')
            if tcW is None: tcW=OxmlElement('w:tcW'); tcPr.append(tcW)
            tcW.set(qn('w:w'),str(w)); tcW.set(qn('w:type'),'dxa')
            set_cell_margins(cell); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

def set_font(run, size=11, bold=False, color=INK, italic=False):
    run.font.name='Calibri'; run._element.rPr.rFonts.set(qn('w:ascii'),'Calibri'); run._element.rPr.rFonts.set(qn('w:hAnsi'),'Calibri')
    run.font.size=Pt(size); run.bold=bold; run.italic=italic; run.font.color.rgb=RGBColor.from_string(color)

def add_page_number(paragraph):
    paragraph.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    r=paragraph.add_run('Page '); set_font(r,9, color='666666')
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); paragraph._p.append(fld)

doc=Document()
sec=doc.sections[0]
sec.top_margin=Inches(1); sec.bottom_margin=Inches(1); sec.left_margin=Inches(1); sec.right_margin=Inches(1)
sec.header_distance=Inches(.492); sec.footer_distance=Inches(.492)

styles=doc.styles
normal=styles['Normal']; normal.font.name='Calibri'; normal._element.rPr.rFonts.set(qn('w:ascii'),'Calibri'); normal._element.rPr.rFonts.set(qn('w:hAnsi'),'Calibri'); normal.font.size=Pt(11); normal.font.color.rgb=RGBColor.from_string(INK); normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.line_spacing=1.1
for name,size,color,before,after in [('Heading 1',16,MID,16,8),('Heading 2',13,MID,12,6),('Heading 3',12,BLUE,8,4)]:
    s=styles[name]; s.font.name='Calibri'; s._element.rPr.rFonts.set(qn('w:ascii'),'Calibri'); s._element.rPr.rFonts.set(qn('w:hAnsi'),'Calibri'); s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor.from_string(color); s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after); s.paragraph_format.keep_with_next=True

footer=sec.footer.paragraphs[0]; add_page_number(footer)
header=sec.header.paragraphs[0]; header.alignment=WD_ALIGN_PARAGRAPH.RIGHT; r=header.add_run('Smart Asset Management System - Updated Proposal'); set_font(r,9,color='666666')

def p(text='', style=None, align=None, before=None, after=None, bold_prefix=None):
    x=doc.add_paragraph(style=style)
    if align is not None: x.alignment=align
    if before is not None: x.paragraph_format.space_before=Pt(before)
    if after is not None: x.paragraph_format.space_after=Pt(after)
    if bold_prefix and text.startswith(bold_prefix):
        a=x.add_run(bold_prefix); set_font(a,11,True)
        b=x.add_run(text[len(bold_prefix):]); set_font(b,11)
    else: x.add_run(text)
    return x

def bullet(text, level=0):
    x=doc.add_paragraph(style='List Bullet' if level==0 else 'List Bullet 2'); x.paragraph_format.space_after=Pt(4); x.paragraph_format.line_spacing=1.1; x.add_run(text); return x

def num(text):
    x=doc.add_paragraph(style='List Number'); x.paragraph_format.space_after=Pt(4); x.paragraph_format.line_spacing=1.1; x.add_run(text); return x

def heading(text, level=1): return doc.add_paragraph(text, style=f'Heading {level}')

def title(text, subtitle=None):
    x=doc.add_paragraph(); x.alignment=WD_ALIGN_PARAGRAPH.CENTER; x.paragraph_format.space_before=Pt(10); x.paragraph_format.space_after=Pt(8); r=x.add_run(text); set_font(r,24,True,BLUE)
    if subtitle:
        x=doc.add_paragraph(); x.alignment=WD_ALIGN_PARAGRAPH.CENTER; x.paragraph_format.space_after=Pt(18); r=x.add_run(subtitle); set_font(r,13,False,'555555')

def simple_table(headers, rows, widths):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.LEFT; t.style='Table Grid'; set_table_widths(t,widths)
    for c,h in zip(t.rows[0].cells,headers):
        set_cell_shading(c,PALE); q=c.paragraphs[0]; q.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=q.add_run(h); set_font(r,10,True,BLUE)
    for row in rows:
        cells=t.add_row().cells
        for c,val in zip(cells,row):
            q=c.paragraphs[0]; q.paragraph_format.space_after=Pt(0); r=q.add_run(str(val)); set_font(r,9.5)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)
    return t

# Cover
for _ in range(4): doc.add_paragraph()
title('MACHINE-LEARNING-ENABLED SMART ASSET MANAGEMENT SYSTEM', 'Predictive Maintenance and Department-Based Laptop Allocation')
for _ in range(3): doc.add_paragraph()
for line in ['Project Proposal (Revised)', 'Kisotu Samwel Lemayian', 'Admission Number: 152170', 'ICS 4B', '', 'Supervisor: Mr. Kevin Ouma', '', 'An Informatics and Computer Science Project Proposal Submitted to the School of Computing and Engineering Sciences in Partial Fulfilment of the Requirements for the Award of a Bachelor of Science in Informatics and Computer Science', '', 'Strathmore University', 'Nairobi, Kenya', 'August 2026']:
    x=doc.add_paragraph(); x.alignment=WD_ALIGN_PARAGRAPH.CENTER; x.paragraph_format.space_after=Pt(8); r=x.add_run(line); set_font(r,11, bool(line in ['Kisotu Samwel Lemayian','Admission Number: 152170','ICS 4B']))
doc.add_page_break()

# Abstract
title('Abstract')
p('This revised proposal presents a web-based Smart Asset Management System for managing organisational ICT assets, with the machine-learning scope focused on predictive maintenance of laptops. The primary model will estimate whether a laptop is likely to require maintenance within the next 30 days using asset age, battery health, battery-cycle count, repair history, fault reports, warranty status, average operating temperature, operating hours and days since previous maintenance. The system will be built with PHP, MySQL, HTML, CSS, JavaScript and XAMPP, while a Python Flask API will expose the trained prediction model to the web application. Where confidential organisational records are unavailable, a documented synthetic dataset will be generated from realistic laptop-maintenance ranges and causal rules. Random Forest will be evaluated against Logistic Regression as a baseline. The project will also provide rule-based lifecycle status, anomaly flags and department-based laptop recommendations; these are decision-support functions, not additional machine-learning models. Google Maps, if used, will visualise authorised location records rather than track devices. Evaluation will include model performance, functional tests, task completion, record completeness, allocation suitability and user feedback. The system aims to improve asset visibility, preventive maintenance planning and the suitability of laptops assigned to departments.')
p('Keywords: ICT asset management; predictive maintenance; laptop allocation; Random Forest; Logistic Regression; synthetic dataset; Flask API.', after=10)
doc.add_page_break()

title('Table of Contents')
for line in ['Abstract', 'Chapter 1: Introduction', 'Chapter 2: Literature Review and Conceptual Framework', 'Chapter 3: Methodology and Evaluation', 'Appendix A: Lecturer Clarifications and Proposed Solutions', 'Appendix B: Synthetic Dataset Specification', 'References']:
    p(line, after=7)
doc.add_page_break()

heading('Chapter 1: Introduction')
heading('1.1 Background',2)
p('ICT-dependent organisations rely on laptops to support administration, finance, communication, teaching, software development and data analysis. Asset registers are often maintained in spreadsheets or separate records, making it difficult to establish who holds a device, its condition, maintenance history and suitability for the work it supports. Reactive maintenance can leave users without a working device and provides little basis for planning repairs or replacements.')
p('This project proposes a centralised web application for registering, assigning, maintaining and reporting on ICT assets. Its intelligent component is deliberately narrow: it predicts a near-term maintenance requirement for laptops so that the ICT team can inspect or service a device before a reported failure interrupts work.')
heading('1.2 Problem Statement',2)
p('Manual or fragmented ICT asset records can be incomplete, outdated and difficult to search. They make it hard to identify laptops with recurring faults or declining condition, resulting in reactive repairs, avoidable downtime and allocation of devices that do not match departmental needs. Existing inventory tools are useful for record keeping but often do not provide a transparent, locally focused preventive-maintenance decision aid.')
heading('1.3 Objectives',2)
p('General objective: To design and develop a secure web-based Smart Asset Management System that improves laptop asset accountability, preventive maintenance planning and department-appropriate allocation.')
p('Specific objectives:', after=4)
for s in ['To analyse current ICT laptop asset-management workflows and challenges.', 'To design and implement a secure web system for laptop registration, assignment, maintenance records, reporting and role-based access.', 'To develop and validate a predictive-maintenance classifier that predicts whether a laptop requires maintenance within the next 30 days.', 'To provide transparent rule-based lifecycle status, anomaly flags and department-based laptop recommendations.', 'To evaluate model quality, system functionality, usability and asset-management effectiveness.']: bullet(s)
heading('1.4 Research Questions',2)
for s in ['How are laptop asset records, assignment and maintenance currently managed?', 'Can laptop condition and maintenance-history data predict a maintenance requirement within 30 days?', 'How accurately does Random Forest perform compared with a Logistic Regression baseline?', 'How can a web system improve maintenance planning, record visibility and the allocation of suitable laptops to departments?']: num(s)
heading('1.5 Scope and Delimitations',2)
p('The project focuses on laptops as the primary predictive-maintenance asset category. Core system functions include asset registration, assignment, maintenance logging, user roles, reporting and alerts. The ML task is one binary classification problem only. Lifecycle classification, anomaly flags and laptop allocation recommendations are rule-based and are included to make the system useful without claiming additional ML models.')
p('The project excludes automatic hardware telemetry collection, procurement, depreciation accounting, advanced hardware diagnostics, continuous GPS tracking and live device tracking by Google Maps. Google Maps may only display a location that has been manually recorded or supplied through an authorised check-in or approved location source.')
heading('1.6 Significance',2)
p('The system supports a shift from a purely reactive maintenance process to a documented, risk-informed process. It also gives administrators a consistent way to select laptop specifications that fit departmental workloads, improving the fit between available devices and user needs.')
doc.add_page_break()

heading('Chapter 2: Literature Review and Conceptual Framework')
heading('2.1 ICT Asset Management and Predictive Maintenance',2)
p('ICT asset management concerns the identification, assignment, maintenance and retirement of technology assets. ISO/IEC 19770-1 describes requirements for IT asset-management systems, while ISO 55000 provides general asset-management principles. Predictive maintenance uses historical condition and event data to estimate a future maintenance need, allowing an organisation to plan intervention before an operational failure (Carvalho et al., 2019; Jardine et al., 2006).')
heading('2.2 Focused Gap',2)
p('Tools such as spreadsheet registers, Snipe-IT and GLPI support inventories, assignments and reports. The project does not attempt to replace complete enterprise IT-service-management suites. Instead, it investigates a focused gap: integrating a transparent laptop-maintenance risk prediction into a simple asset-management workflow and combining it with department-specific allocation guidance.')
heading('2.3 Conceptual Framework',2)
p('The system has four connected layers. The web layer captures and displays asset data; MySQL stores authorised records; the prediction service scores a laptop from the defined features; and the decision-support layer presents a risk result together with clear actions. The model is trained offline and versioned before deployment. It is not trained automatically from live user activity.')
simple_table(['Layer','Function','Output'],[
['Web application (PHP)','Asset registration, assignment, maintenance log, authorised check-in, roles and reports','Validated asset records and administrator actions'],
['Database (MySQL)','Stores laptops, users, departments, maintenance and assignment history','Structured records for reporting and scoring'],
['ML service (Python/Flask)','Preprocesses inputs and runs the selected Random Forest classifier','Maintenance probability and high/medium/low risk label'],
['Decision-support rules','Applies lifecycle, anomaly and allocation rules','Replacement status, flags and recommended laptop profile']], [1900,3500,3960])
heading('2.4 What Is in the Web System',2)
for s in ['Asset registration: serial number, model, purchase date, specifications, warranty and condition.', 'Assignment and return: assigned employee, department, issue date, current office/location record and handover history.', 'Maintenance: fault tickets, repair count, service dates, service outcomes and maintenance alerts.', 'Decision support: 30-day maintenance risk, lifecycle label, anomaly flags and department-match recommendation.', 'Security: authenticated users, role-based permissions and audit entries.']: bullet(s)
heading('2.5 Laptop Allocation Profiles',2)
p('Allocation is a transparent recommendation, not an ML prediction. An administrator enters the recipient department and the system lists available laptops that meet the minimum profile, have an acceptable lifecycle label and are not high maintenance risk.')
simple_table(['Department','Minimum recommended profile','Reason'],[
['Software development / ICT','Core i7 or Ryzen 7; 16 GB RAM; 512 GB SSD','Development tools, virtual machines and multitasking'],
['Data analysis / research','Core i7 or Ryzen 7; 16-32 GB RAM; 512 GB SSD','Analysis software and larger datasets'],
['Design / media','Core i7 or Ryzen 7; 16-32 GB RAM; 512 GB SSD; dedicated graphics where needed','Graphics and media editing'],
['Finance / administration / HR','Core i5 or Ryzen 5; 8-16 GB RAM; 256-512 GB SSD','Office suites, information systems and routine multitasking'],
['Management','Core i5 or better; 16 GB RAM; 512 GB SSD; strong battery life','Portable productivity and meetings'],
['Reception / data entry','Core i3/i5; 8 GB RAM; 256 GB SSD','Browser, email and basic transaction work']], [2100,4200,3060])
doc.add_page_break()

heading('Chapter 3: Methodology and Evaluation')
heading('3.1 Research and Development Approach',2)
p('The project uses an applied design-science approach: identify an asset-management problem, build an artefact, demonstrate it with controlled data and evaluate its technical and practical usefulness. The software will be developed iteratively using OOAD and Scrum for One. Each sprint will produce a testable increment, beginning with the asset register and maintenance log, followed by the model API, decision-support screens and evaluation.')
heading('3.2 Primary Machine-Learning Task',2)
p('The primary ML task is binary classification: predict whether an individual laptop is likely to require maintenance in the next 30 days. A positive result means that the asset should be reviewed by ICT; it is not a claim that the laptop has already failed. This action-oriented definition makes the prediction measurable and aligned to a maintenance workflow.')
heading('3.3 Dataset, Source and Synthetic-Data Justification',2)
p('The project will use a structured synthetic laptop-maintenance dataset because real organisational repair and device-condition records may be confidential or unavailable. It will be generated in Python using documented value ranges and conditional rules informed by standard laptop-maintenance practice and the variables discussed in predictive-maintenance literature. No personal employee data will be used.')
p('Each record represents a laptop assessment at a point in time. Values will be sampled within plausible ranges, then the target label will be assigned with a documented risk function: devices with greater age, more repairs/fault reports, reduced battery health, high temperature, many operating hours, expired warranty and a long period since service receive a higher probability of a positive label. Controlled random variation will prevent a perfectly deterministic dataset. The generator, random seed and data dictionary will be retained so the experiment can be reproduced.')
heading('3.4 Target Variable and Features',2)
simple_table(['Type','Variables'],[
['Target','maintenance_required_within_30_days (1 = yes; 0 = no)'],
['Predictor features','age_months; battery_health_percent; battery_cycle_count; repair_count_12_months; fault_reports_90_days; days_since_last_maintenance; average_operating_temperature; operating_hours; warranty_active; storage_free_percent; RAM utilisation band; department'],
['Excluded from model','asset ID, serial number, employee name, phone number and exact location, because they do not represent a device-condition cause and may introduce privacy or leakage concerns']], [1800,7560])
heading('3.5 Preprocessing and Data Split',2)
p('Records will be checked for duplicates, invalid ranges and missing values. Numeric missing values will be imputed using the training-set median; categorical values will use a defined “unknown” category. Categorical department values will be one-hot encoded. The data will be split using stratified sampling to preserve the positive/negative class ratio: 70% training, 15% validation and 15% final test data. The test set will remain untouched until final evaluation. The generator will avoid copying target information directly into a predictor.')
heading('3.6 Models and Comparison',2)
p('Logistic Regression will be the baseline classifier because it is simple, interpretable and establishes whether a basic linear model is adequate. Random Forest will be the proposed model because it can learn nonlinear relationships and interactions among device condition, repair and usage features. Hyperparameters will be selected with validation data. Both models will use identical preprocessing and the same held-out test set. The model selected for the web system will be the one with stronger recall and F1-score while maintaining acceptable precision and an understandable decision threshold.')
heading('3.7 Model Evaluation and Acceptance Criteria',2)
simple_table(['Measure','Use in this project','Acceptance criterion'],[
['Recall','Proportion of true maintenance cases identified. Important because missed cases can cause downtime.','At least 0.80 on the test set'],
['Precision','Proportion of flagged laptops that genuinely need maintenance. Limits unnecessary inspections.','At least 0.65 on the test set'],
['F1-score','Balances precision and recall for the positive class.','At least 0.75 and higher than baseline'],
['ROC-AUC and confusion matrix','Shows discrimination and types of error.','Reported for both models; supports threshold selection'],
['Baseline comparison','Tests whether Random Forest improves on Logistic Regression.','Random Forest must equal or exceed baseline F1 and recall']], [1700,4700,2960])
heading('3.8 Rule-Based Support Functions',2)
p('Lifecycle status is a measurable rule-based classification: healthy (within age/warranty thresholds and low risk); monitor (moderate risk or declining health); service due (high predicted risk or repeated faults); replacement review (for example, expired warranty plus age threshold and recurring repairs); and retired (authorised disposal/retirement record). This is explicitly not a separate lifecycle-prediction ML model.')
p('An anomaly is a record or event inconsistent with approved asset-management rules. Examples are an asset assigned to two people simultaneously, a location change without an authorised transfer/check-in, repeated repair tickets within 90 days, a retired asset marked active, or an unavailable device allocated to a user. These flags are generated with transparent rules, logged for review and never automatically punish a user.')
heading('3.9 System Evaluation',2)
p('Functional tests will confirm that users can register, search, assign, return, update and report on laptops; that authorisation protects restricted actions; and that the Flask endpoint validates input and returns the correct model result. System effectiveness will be assessed in a controlled usability evaluation with representative administrator tasks and a short questionnaire.')
simple_table(['Measure','How it will be assessed'],[
['Asset-record completeness','Percentage of sampled records containing required owner/department, serial number, condition, warranty and maintenance fields.'],
['Asset retrieval efficiency','Median time to locate a specified laptop, its assignment and latest maintenance status.'],
['Preventive-maintenance coverage','Percentage of high-risk test assets flagged for review before their simulated maintenance event.'],
['Allocation suitability','Percentage of recommended laptops that meet the stated department profile and pass availability/lifecycle checks.'],
['Usability and satisfaction','Task-completion rate, task time and 5-point user feedback on clarity, usefulness and trust.'],
['Comparison','Compare task results using a spreadsheet-style baseline workflow versus the prototype where feasible.']], [2500,6860])
heading('3.10 Ethical and Practical Considerations',2)
p('Synthetic data avoids disclosure of confidential organisational records. The model output is a maintenance recommendation for administrator review, not an automatic repair, disposal or disciplinary decision. System data will be protected through access controls and audit logs. Any future use of location information will require organisational authorisation and a clear retention policy.')
doc.add_page_break()

heading('Appendix A: Lecturer Clarifications and Proposed Solutions')
qa=[
('1. Machine-learning scope','Predictive maintenance is the sole ML task. The model predicts maintenance_required_within_30_days for laptops. Lifecycle status, allocation and anomaly functions are rule-based decision support, which keeps the project feasible and removes the original over-broad ML scope.'),
('2. Dataset','A documented synthetic laptop-maintenance dataset will be generated in Python from plausible ranges and conditional risk rules. The target is whether maintenance is required within 30 days. The features, label logic, random seed and data dictionary are specified in Appendix B and Section 3.4.'),
('3. Model inconsistency','Logistic Regression is the baseline. Random Forest is the proposed final classifier. They will be trained on the same prepared data and compared with recall, precision, F1-score, ROC-AUC and a confusion matrix. No separate lifecycle model is proposed.'),
('4. Model evaluation','Data will be stratified into 70% training, 15% validation and 15% test sets. The test set is used once for final reporting. An acceptable model has test recall of at least 0.80, F1-score at least 0.75, precision at least 0.65 and performance that meets or improves on the Logistic Regression baseline.'),
('5. Anomaly detection','An anomaly is an inconsistent asset-management event: duplicate assignment, unauthorised location change, repeated repair tickets, retired asset shown as active, or allocation of an unavailable device. Rule thresholds are explicit and alerts require administrator review.'),
('6. Lifecycle and allocation','Lifecycle is a measurable status label based on age, warranty, battery health, repair history and maintenance risk. Allocation recommends an available, low-risk laptop whose specifications meet the minimum profile for the recipient department. Allocation quality is measured by profile compliance and user/task feedback.'),
('7. Google-based tracking','Google Maps does not track laptops. It can only visualise an authorised stored location, such as a department office, check-in record or coordinates provided by a separate approved source. The proposal makes this distinction and excludes automatic GPS tracking.'),
('8. System effectiveness','Beyond functional tests, the system will measure asset-record completeness, asset retrieval time, preventive-maintenance coverage, allocation suitability, task-completion rate and user satisfaction. Where feasible, these will be compared with a spreadsheet-style baseline process.')]
for q,a in qa:
    heading(q,2); p(a)
doc.add_page_break()

heading('Appendix B: Synthetic Dataset Specification')
p('Suggested dataset size: 1,200 laptop assessment records. The number may be adjusted after exploratory analysis, but the class distribution and generation logic must be reported. The dataset will be used only for development and academic demonstration; it does not represent observed organisational failure rates.')
simple_table(['Field','Example range / coding','Rationale'],[
['age_months','6-84 months','Older devices may have higher maintenance risk.'],
['battery_health_percent','35-100','Low health indicates degradation.'],
['battery_cycle_count','50-1,200','High cycles indicate battery wear.'],
['repair_count_12_months','0-6','Repeated repairs increase risk.'],
['fault_reports_90_days','0-5','Recent faults indicate emerging condition issues.'],
['Days since service','0-730','Longer unserviced periods can increase risk.'],
['Average temperature','35-95 degrees C','Persistent high temperature can indicate thermal stress.'],
['operating_hours','100-20,000','Usage exposure.'],
['warranty_active','0/1','Warranty helps maintenance/replacement planning.'],
['department','categorical','Supports allocation context; does not replace condition features.'],
['Maintenance needed (30 days)','0/1','Target label created from the documented risk function plus noise.']], [2500,2500,4360])
heading('References')
refs=['Carvalho, T. P., Soares, F. A. A. M. N., Vita, R., Francisco, R. D. P., Basto, J. P., & Alcala, S. G. S. (2019). A systematic literature review of machine learning methods applied to predictive maintenance. Computers & Industrial Engineering, 137, 106024. https://doi.org/10.1016/j.cie.2019.106024', 'Géron, A. (2022). Hands-on machine learning with Scikit-Learn, Keras, and TensorFlow (3rd ed.). O\'Reilly Media.', 'International Organization for Standardization. (2014). ISO 55000:2014 Asset management - Overview, principles and terminology.', 'ISO/IEC. (2017). ISO/IEC 19770-1:2017 Information technology - IT asset management - Part 1: IT asset management systems - Requirements.', 'Jardine, A. K. S., Lin, D., & Banjevic, D. (2006). A review on machinery diagnostics and prognostics implementing condition-based maintenance. Mechanical Systems and Signal Processing, 20(7), 1483-1510. https://doi.org/10.1016/j.ymssp.2005.09.012', 'Pedregosa, F., et al. (2011). Scikit-learn: Machine learning in Python. Journal of Machine Learning Research, 12, 2825-2830.']
for r in refs: p(r, after=8)

doc.core_properties.title='Smart Asset Management System - Updated Proposal'
doc.core_properties.author='Kisotu Samwel Lemayian'
doc.save(OUT)
print(OUT)
